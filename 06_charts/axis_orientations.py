# axis_orientations.py

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout


def create_chart(title, sheet, x_orientation=None, y_orientation=None):
    chart = BarChart()
    chart.title = title
    chart.x_axis.title = "Book Types"
    chart.y_axis.title = "Prices"

    data = Reference(worksheet=sheet,
                     min_row=1,
                     max_row=10,
                     min_col=2,
                     max_col=3)
    chart.add_data(data, titles_from_data=True)

    if x_orientation:
        chart.x_axis.scaling.orientation = x_orientation
    if y_orientation:
        chart.y_axis.scaling.orientation = y_orientation

    return chart


def main(filename):
    workbook = Workbook()
    sheet = workbook.active

    # Add data to spreadsheet
    data_rows = [
        ["Book", "Kindle", "Paperback"],
        [1, 9.99, 15.99],
        [2, 9.99, 25.99],
        [3, 9.99, 25.99],
        [4, 4.99, 29.99],
        [5, 14.99, 39.99],
    ]

    for row in data_rows:
        sheet.append(row)

    # Create the bar charts
    sheet.add_chart(create_chart("Defaults", sheet), "D1")
    sheet.add_chart(create_chart("Flip X", sheet,
                                 x_orientation="maxMin",
                                 y_orientation="minMax"), "J1")
    sheet.add_chart(create_chart("Flip Y", sheet,
                                 x_orientation="minMax",
                                 y_orientation="maxMin"), "D15")
    sheet.add_chart(create_chart("Flip Both", sheet,
                                 x_orientation="maxMin",
                                 y_orientation="maxMin"), "J15")

    workbook.save(filename)


if __name__ == "__main__":
    main("axis_orientations.xlsx")