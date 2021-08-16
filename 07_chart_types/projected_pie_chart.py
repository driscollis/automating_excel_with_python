# projected_pie_chart.py

from copy import deepcopy

from openpyxl import Workbook
from openpyxl.chart import ProjectedPieChart, Reference


def main(filename):
    workbook = Workbook()
    sheet = workbook.active

    data = [
        ['Page', 'Views'],
        ['Search', 95],
        ['Products', 4],
        ['Offers', 0.5],
        ['Sales', 0.5],
    ]

    for row in data:
        sheet.append(row)

    projected_pie = ProjectedPieChart()
    projected_pie.type = "pie"
    projected_pie.splitType = "val" # split by value
    labels = Reference(sheet, min_col=1, min_row=2, max_row=5)
    data = Reference(sheet, min_col=2, min_row=1, max_row=5)
    projected_pie.add_data(data, titles_from_data=True)
    projected_pie.set_categories(labels)
    sheet.add_chart(projected_pie, "E2")

    projected_bar = deepcopy(projected_pie)
    projected_bar.type = "bar"
    projected_bar.splitType = 'pos' # split by position
    sheet.add_chart(projected_bar, "E19")

    workbook.save(filename)


if __name__ == "__main__":
    main("project_pie_chart.xlsx")
