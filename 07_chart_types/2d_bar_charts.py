# 2d_bar_charts.py

from copy import deepcopy
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


def create_excel_data(sheet):
    data_rows = [
        ["Book", "Kindle", "Paperback"],
        [1, 9.99, 25.99],
        [2, 9.99, 25.99],
        [3, 9.99, 25.99],
        [4, 4.99, 29.99],
        [5, 4.99, 29.99],
        [6, 24.99, 29.99],
        [7, 24.99, 65.00],
        [8, 24.99, 69.00],
        [9, 24.99, 69.00],
    ]

    for row in data_rows:
        sheet.append(row)


def create_bar_charts(sheet):
    bar_chart = BarChart()

    data = Reference(worksheet=sheet,
                     min_row=1,
                     max_row=10,
                     min_col=2,
                     max_col=3)
    bar_chart.add_data(data, titles_from_data=True)
    sheet.add_chart(bar_chart, "A11")

    # Add more charts!
    chart2 = deepcopy(bar_chart)
    chart2.style = 11
    chart2.type = "bar"
    chart2.title = "Horizontal Bar Chart"
    sheet.add_chart(chart2, "G11")

    chart3 = deepcopy(bar_chart)
    chart3.type = "col"
    chart3.style = 12
    chart3.grouping = "stacked"
    chart3.overlap = 100
    chart3.title = 'Stacked Chart'
    sheet.add_chart(chart3, "A26")

    chart4 = deepcopy(bar_chart)
    chart4.type = "bar"
    chart4.style = 13
    chart4.grouping = "percentStacked"
    chart4.overlap = 100
    chart4.title = 'Percent Stacked Chart'
    sheet.add_chart(chart4, "G26")


def main():
    workbook = Workbook()
    sheet = workbook.active
    create_excel_data(sheet)
    create_bar_charts(sheet)
    workbook.save("2d_bar_charts.xlsx")


if __name__ == "__main__":
    main()