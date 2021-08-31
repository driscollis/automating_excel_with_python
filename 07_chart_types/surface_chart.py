# surface_chart.py

from copy import deepcopy

from openpyxl import Workbook
from openpyxl.chart import SurfaceChart, SurfaceChart3D
from openpyxl.chart import Reference


def main(filename):
    workbook = Workbook()
    sheet = workbook.active
    
    data = [
        [None, 10, 20, 30, 40, 50,],
        [0.1, 15, 65, 105, 65, 15,],
        [0.2, 35, 105, 170, 105, 35,],
        [0.3, 55, 135, 215, 135, 55,],
        [0.4, 75, 155, 240, 155, 75,],
        [0.5, 80, 190, 245, 190, 80,],
        [0.6, 75, 155, 240, 155, 75,],
        [0.7, 55, 135, 215, 135, 55,],
        [0.8, 35, 105, 170, 105, 35,],
        [0.9, 15, 65, 105, 65, 15],
    ]
    
    for row in data:
        sheet.append(row)


    chart1 = SurfaceChart()
    ref = Reference(sheet, min_col=2, max_col=6, min_row=1, max_row=10)
    labels = Reference(sheet, min_col=1, min_row=2, max_row=10)
    chart1.add_data(ref, titles_from_data=True)
    chart1.set_categories(labels)
    chart1.title = "Contour"
    
    sheet.add_chart(chart1, "A12")
    
    # wireframe
    chart2 = deepcopy(chart1)
    chart2.wireframe = True
    chart2.title = "2D Wireframe"
    
    sheet.add_chart(chart2, "G12")
    
    # 3D Surface
    chart3 = SurfaceChart3D()
    chart3.add_data(ref, titles_from_data=True)
    chart3.set_categories(labels)
    chart3.title = "Surface"
    
    sheet.add_chart(chart3, "A29")
    
    chart4 = deepcopy(chart3)
    chart4.wireframe = True
    chart4.title = "3D Wireframe"
    
    sheet.add_chart(chart4, "G29")
    
    workbook.save(filename)

if __name__ == "__main__":
    main("surface_chart.xlsx")
